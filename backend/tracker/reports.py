"""Profit / loss report data plus PDF and Excel exports."""
import io
from datetime import date
from decimal import Decimal

from django.db.models import Sum
from django.db.models.functions import TruncMonth
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Business, Expense, Income

ZERO = Decimal('0')


def build_report(user, start=None, end=None, business_id=None):
    businesses = Business.objects.filter(owner=user)
    if business_id:
        businesses = businesses.filter(pk=business_id)
    expenses = Expense.objects.filter(business__in=businesses)
    incomes = Income.objects.filter(business__in=businesses)
    if start:
        expenses, incomes = expenses.filter(date__gte=start), incomes.filter(date__gte=start)
    if end:
        expenses, incomes = expenses.filter(date__lte=end), incomes.filter(date__lte=end)

    total_expense = expenses.aggregate(s=Sum('amount'))['s'] or ZERO
    total_income = incomes.aggregate(s=Sum('amount'))['s'] or ZERO

    # Month-by-month series with a running (cumulative) profit
    months = {}
    for qs, key in ((expenses, 'expense'), (incomes, 'income')):
        for row in qs.annotate(m=TruncMonth('date')).values('m').annotate(s=Sum('amount')):
            month = row['m'].strftime('%Y-%m')
            months.setdefault(month, {'income': ZERO, 'expense': ZERO})[key] = row['s']
    running = ZERO
    by_month = []
    for month in sorted(months):
        income, expense = months[month]['income'], months[month]['expense']
        running += income - expense
        by_month.append({'month': month, 'income': income, 'expense': expense,
                         'profit': income - expense, 'cumulative_profit': running})

    labels = dict(Expense.Category.choices)
    expense_by_category = [
        {'category': row['category'], 'label': labels[row['category']], 'amount': row['s']}
        for row in expenses.values('category').annotate(s=Sum('amount')).order_by('-s')
    ]
    source_labels = dict(Income.Source.choices)
    income_by_source = [
        {'source': row['source'], 'label': source_labels[row['source']], 'amount': row['s']}
        for row in incomes.values('source').annotate(s=Sum('amount')).order_by('-s')
    ]

    exp_by_biz = dict(expenses.values_list('business').annotate(s=Sum('amount')))
    inc_by_biz = dict(incomes.values_list('business').annotate(s=Sum('amount')))
    by_business = []
    for b in businesses.order_by('-is_active', 'name'):
        inc, exp = inc_by_biz.get(b.pk, ZERO), exp_by_biz.get(b.pk, ZERO)
        if inc or exp or business_id:
            by_business.append({'id': b.pk, 'name': b.name, 'business_type': b.business_type,
                                'income': inc, 'expense': exp, 'profit': inc - exp})

    return {
        'start': start, 'end': end,
        'business': businesses.first().name if business_id and businesses.exists() else None,
        'totals': {'income': total_income, 'expense': total_expense, 'profit': total_income - total_expense},
        'by_month': by_month,
        'expense_by_category': expense_by_category,
        'income_by_source': income_by_source,
        'by_business': by_business,
        'expenses': expenses.select_related('business').order_by('date'),
        'incomes': incomes.select_related('business').order_by('date'),
    }


def _period_label(report):
    start = report['start'].strftime('%d %b %Y') if report['start'] else 'Beginning'
    end = report['end'].strftime('%d %b %Y') if report['end'] else date.today().strftime('%d %b %Y')
    return f'{start} – {end}'


def _rs(value):
    return f'Rs {value:,.0f}'


# ---------- Excel ----------

HEADER_FILL = PatternFill('solid', fgColor='1E5E2F')
HEADER_FONT = Font(bold=True, color='FFFFFF')


def _sheet(wb, title, headers, rows, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for row in rows:
        ws.append(row)
    for col in money_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = '#,##0'
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [10])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 12), 50)
    ws.freeze_panes = 'A2'
    return ws


def export_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = 'Agri Shop – Profit & Loss Report'
    ws['A1'].font = Font(bold=True, size=16, color='1E5E2F')
    ws['A2'] = f"Business: {report['business'] or 'All businesses'}"
    ws['A3'] = f'Period: {_period_label(report)}'
    rows = [('Total income', report['totals']['income']), ('Total expense', report['totals']['expense']),
            ('Profit / Loss', report['totals']['profit'])]
    for i, (label, value) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=float(value))
        cell.number_format = '#,##0'
        cell.alignment = Alignment(horizontal='right')
    ws['B7'].font = Font(bold=True, color='1E5E2F' if report['totals']['profit'] >= 0 else 'B91C1C')
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 22, 18

    _sheet(wb, 'By business', ['Business', 'Income', 'Expense', 'Profit'],
           [(b['name'], float(b['income']), float(b['expense']), float(b['profit'])) for b in report['by_business']],
           money_cols=(2, 3, 4))
    _sheet(wb, 'Monthly', ['Month', 'Income', 'Expense', 'Profit', 'Cumulative profit'],
           [(m['month'], float(m['income']), float(m['expense']), float(m['profit']), float(m['cumulative_profit']))
            for m in report['by_month']], money_cols=(2, 3, 4, 5))
    _sheet(wb, 'Expenses', ['Date', 'Business', 'Category', 'Amount', 'Notes'],
           [(e.date, e.business.name, e.get_category_display(), float(e.amount), e.notes) for e in report['expenses']],
           money_cols=(4,))
    _sheet(wb, 'Income', ['Date', 'Business', 'Source', 'Amount', 'Notes'],
           [(i.date, i.business.name, i.get_source_display(), float(i.amount), i.notes) for i in report['incomes']],
           money_cols=(4,))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------

GREEN = colors.HexColor('#1E5E2F')
LIGHT = colors.HexColor('#F0F9F1')


def _table(data, col_widths, money_from=1):
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), GREEN),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (money_from, 0), (-1, -1), 'RIGHT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, LIGHT]),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#D6D3D1')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    return table


def export_pdf(report):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm, title='Profit & Loss Report')
    styles = getSampleStyleSheet()
    styles['Title'].textColor = GREEN
    styles['Heading2'].textColor = GREEN
    totals = report['totals']
    profit_color = '#1E5E2F' if totals['profit'] >= 0 else '#B91C1C'
    width = doc.width

    story = [
        Paragraph('Agri Shop – Profit &amp; Loss Report', styles['Title']),
        Paragraph(f"<b>Business:</b> {report['business'] or 'All businesses'} &nbsp;&nbsp; "
                  f"<b>Period:</b> {_period_label(report)}", styles['Normal']),
        Spacer(1, 8 * mm),
        _table([['Total income', 'Total expense', 'Profit / Loss'],
                [_rs(totals['income']), _rs(totals['expense']),
                 Paragraph(f'<font color="{profit_color}"><b>{_rs(totals["profit"])}</b></font>', styles['Normal'])]],
               [width / 3] * 3, money_from=3),
        Spacer(1, 6 * mm),
    ]
    if report['by_business']:
        story += [Paragraph('By business', styles['Heading2']),
                  _table([['Business', 'Income', 'Expense', 'Profit']] +
                         [[b['name'], _rs(b['income']), _rs(b['expense']), _rs(b['profit'])] for b in report['by_business']],
                         [width * 0.4] + [width * 0.2] * 3),
                  Spacer(1, 6 * mm)]
    if report['by_month']:
        story += [Paragraph('Monthly', styles['Heading2']),
                  _table([['Month', 'Income', 'Expense', 'Profit', 'Cumulative']] +
                         [[m['month'], _rs(m['income']), _rs(m['expense']), _rs(m['profit']), _rs(m['cumulative_profit'])]
                          for m in report['by_month']],
                         [width * 0.2] * 5),
                  Spacer(1, 6 * mm)]
    if report['expense_by_category']:
        total = totals['expense'] or 1
        story += [Paragraph('Expenses by category', styles['Heading2']),
                  _table([['Category', 'Amount', 'Share']] +
                         [[c['label'], _rs(c['amount']), f"{c['amount'] * 100 / total:.1f}%"] for c in report['expense_by_category']],
                         [width * 0.5, width * 0.3, width * 0.2])]

    doc.build(story)
    return buffer.getvalue()
